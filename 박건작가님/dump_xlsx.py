import openpyxl
import json

wb = openpyxl.load_workbook(r'C:\Users\passp\Desktop\univercity\4-2\머신아트랩\박건작가님\1차부품정리리스트.xlsx')
out = {}
for name in wb.sheetnames:
    s = wb[name]
    data = []
    for row in s.iter_rows(values_only=True):
        if any(v is not None for v in row):
            data.append([str(v) if v is not None else "" for v in row])
    out[name] = data

with open(r'C:\Users\passp\Desktop\univercity\4-2\머신아트랩\박건작가님\xlsx_parsed.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=2)
print("Done")
