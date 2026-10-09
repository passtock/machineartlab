# -*- coding: utf-8 -*-
"""
[머신아트랩 - 박건 작가님] 인터랙티브 시스템 부품 리스트 및 배선 가이드 docx 생성기
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_hyperlink(paragraph, url, text, color="0066CC", underline=True, bold=True):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    rPr.append(rFonts)
    
    c = OxmlElement('w:color')
    c.set(qn('w:val'), color)
    rPr.append(c)
    
    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)
        
    if bold:
        b = OxmlElement('w:b')
        rPr.append(b)
        
    new_run.append(rPr)
    text_elm = OxmlElement('w:t')
    text_elm.text = text
    new_run.append(text_elm)
    
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def create_document():
    doc = docx.Document()

    # 페이지 여백 설정 (상하 20mm, 좌우 20mm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 기본 스타일 설정
    style = doc.styles['Normal']
    style.font.name = '맑은 고딕'
    style.font.size = Pt(10)
    style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

    # ---------------------------------------------------------
    # 1. 문서 헤더 / 타이틀
    # ---------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("머신아트랩 (Machine Art Lab) — 프로젝트 기술 문서\n")
    run_sub.font.size = Pt(10.5)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

    run_title = title_p.add_run("[박건 작가님] 인터랙티브 디스플레이 부품 목록 & 배선·납땜 가이드")
    run_title.font.size = Pt(19)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A)

    desc_p = doc.add_paragraph()
    desc_p.paragraph_format.space_before = Pt(4)
    desc_p.paragraph_format.space_after = Pt(16)
    run_desc = desc_p.add_run(
        "본 문서는 ESP32 스마트 디스플레이(UEDX80480043E-WB-B), MAX7219 8자리 7세그먼트, 5kg 로드셀(HX711)을 "
        "연동하기 위한 프로토타이핑/납땜/배선 자재(브레드보드, 핀헤더, 점퍼 케이블, 인두기) 목록과 하드웨어 1:1 결선도를 정리한 공식 사양서입니다."
    )
    run_desc.font.size = Pt(9.5)
    run_desc.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # 구분선
    line_p = doc.add_paragraph()
    line_p.paragraph_format.space_before = Pt(0)
    line_p.paragraph_format.space_after = Pt(12)
    run_line = line_p.add_run("―" * 58)
    run_line.font.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)

    # ---------------------------------------------------------
    # 2. 신규 자재: 브레드보드, 핀헤더, 전선, 납땜도구 목록
    # ---------------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. 프로토타이핑 & 배선·납땜 자재 목록 (디바이스마트 실시간 검증)")
    r_h1.font.size = Pt(13)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A)

    table1_data = [
        ("No.", "품목명 및 규격", "수량", "상품번호", "단가(예상)", "핵심 용도 및 연결 방식", "구매 링크"),
        (
            "1",
            "브레드보드 400핀\n(Half Size 빵판)",
            "1개",
            "1328148",
            "935원",
            "전원 분배 허브 (디스플레이 5V/GND를 세그먼트와 HX711로 공유 분배) 및 센서 임시 고정",
            "https://www.devicemart.co.kr/goods/view?no=1328148"
        ),
        (
            "2",
            "2.54mm 1열 40핀 핀헤더\n(수 / Male Straight)",
            "2개",
            "2825",
            "200원",
            "디스플레이 보드 상단 J2 스루홀 및 HX711 모듈 기판 납땜용 (니퍼로 필요한 핀 수만큼 톡 잘라 사용)",
            "https://www.devicemart.co.kr/goods/view?no=2825"
        ),
        (
            "3",
            "2.54mm 1열 40핀 핀소켓\n(암 / Female Straight)",
            "1개",
            "12493",
            "530원",
            "[선택] 디스플레이 보드에 구멍 소켓 형태로 납땜하여 쇼트(합선) 방지 및 슬림 매립용",
            "https://www.devicemart.co.kr/goods/view?no=12493"
        ),
        (
            "4",
            "점퍼 케이블 40P 암-수\n(M/F) 20cm",
            "1세트",
            "1321195",
            "1,100원",
            "디스플레이(암소켓/수핀) ➔ 브레드보드 전원 레일 및 센서 모듈 연결 시 가장 많이 사용되는 필수선",
            "https://www.devicemart.co.kr/goods/view?no=1321195"
        ),
        (
            "5",
            "점퍼 케이블 40P 암-암\n(F/F) 20cm",
            "1세트",
            "1321192",
            "1,100원",
            "디스플레이 보드(수핀 납땜 시) ➔ 세그먼트 디스플레이(수핀) 직결 신호선(DIN/CLK/CS) 연결용",
            "https://www.devicemart.co.kr/goods/view?no=1321192"
        ),
        (
            "6",
            "점퍼 케이블 40P 수-수\n(M/M) 20cm",
            "1세트",
            "1328409",
            "1,100원",
            "브레드보드 내부 전원 레일 연결 및 배선 확장용",
            "https://www.devicemart.co.kr/goods/view?no=1328409"
        ),
        (
            "7",
            "30W 전기인두기 9종 세트\n(가방 포함 납땜 풀세트)",
            "1세트",
            "13237150",
            "약 14,000원",
            "전기인두기 + 거치대 + 실납 + 핀셋 + 흡입기 등 포함. J2 핀헤더 및 로드셀 배선 납땜 작업 풀패키지",
            "https://www.devicemart.co.kr/goods/view?no=13237150"
        )
    ]

    t1 = doc.add_table(rows=len(table1_data), cols=7)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths1 = [Inches(0.4), Inches(1.5), Inches(0.5), Inches(0.8), Inches(0.7), Inches(2.2), Inches(0.8)]

    for row_idx, row in enumerate(t1.rows):
        data_row = table1_data[row_idx]
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths1[col_idx]
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            if row_idx == 0:
                set_cell_background(cell, "0F294A")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(data_row[col_idx])
                run.font.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if col_idx in [0, 2, 3, 4]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                if col_idx == 6: # 링크 열
                    add_hyperlink(p, data_row[col_idx], "상품보기", color="0066CC", underline=True, bold=True)
                else:
                    run = p.add_run(data_row[col_idx])
                    run.font.size = Pt(8.5)
                    if col_idx == 1:
                        run.font.bold = True

    # ---------------------------------------------------------
    # 3. 메인 하드웨어 장치 목록
    # ---------------------------------------------------------
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("2. 전체 인터랙티브 시스템 핵심 하드웨어 장치 목록")
    r_h2.font.size = Pt(13)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A)

    table2_data = [
        ("No.", "부품명", "규격 및 모델명", "수량", "핵심 역할 및 동작 메커니즘"),
        (
            "1",
            "ESP32-S3 4.3인치\n스마트 디스플레이",
            "VIEWE UEDX80480043E-WB-B (V1.1)\n(800x480 RGB IPS, 16MB Flash, 8MB PSRAM)",
            "1개",
            "전체 시스템 메인 제어기 + 사진 화면 출력. 터치 감지 시 SD 카드의 JPEG 사진을 0.1초 만에 화면 전체에 초고속 렌더링하고 세그먼트 애니메이션 총괄 제어"
        ),
        (
            "2",
            "8자리 7세그먼트\n디스플레이 모듈",
            "MAX7219 8-Digit 0.36\" LED Display\n(SPI 3선 직렬 제어, 3.3V/5V 겸용)",
            "1개",
            "관람객이 손을 올렸을 때 '띠리리리' 카지노 룰렛 감속 롤링 애니메이션 재생 후 8자리 행운 번호 고정 표시 (평상시: '--------' 대기)"
        ),
        (
            "3",
            "로드셀 센서 키트\n(손 터치 감지)",
            "5kg 알루미늄 외팔보 로드셀 +\nHX711 24비트 고정밀 ADC 앰프 모듈",
            "1세트",
            "상부 100mm 원판에 관람객이 손을 얹는 순간 발생하는 미세 하중(100g 이상)을 정밀 감지하여 화면 전환 및 룰렛 추첨 트리거 발생"
        ),
        (
            "4",
            "고속 MicroSD 카드",
            "SanDisk Ultra MicroSDHC 32GB\n(Class 10 / UHS-I A1 고속 판독)",
            "1개",
            "디스플레이 보드 TF 슬롯에 장착. '/images' 폴더 내 100여 장 이상의 고화질 JPEG 이미지를 보관하며 고속 로드 지원"
        ),
        (
            "5",
            "시스템 주 전원 장치",
            "정격 출력 DC 5V 3A (15W) Type-C\n어댑터 및 고전류 데이터 케이블",
            "1개",
            "ESP32 프로세서, 800x480 고휘도 백라이트, 8자리 LED 세그먼트의 순간 전력 소모를 안정적으로 감당하여 재부팅/화면 꺼짐 방지"
        ),
        (
            "6",
            "패시브 피에조 부저\n(선택 사항)",
            "아두이노 패시브 부저 모듈\n(수동 주파수 사각파 tone 제어)",
            "1개",
            "숫자 롤링 시 600Hz에서 2400Hz로 피치가 급상승하는 룰렛 효과음 연주 및 최종 확정 '띵-!' 사운드 출력"
        )
    ]

    t2 = doc.add_table(rows=len(table2_data), cols=5)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths2 = [Inches(0.4), Inches(1.6), Inches(2.1), Inches(0.5), Inches(2.3)]

    for row_idx, row in enumerate(t2.rows):
        data_row = table2_data[row_idx]
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths2[col_idx]
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            if row_idx == 0:
                set_cell_background(cell, "0F294A")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(data_row[col_idx])
                run.font.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if col_idx in [0, 3]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(data_row[col_idx])
                run.font.size = Pt(8.5)
                if col_idx == 1:
                    run.font.bold = True

    # ---------------------------------------------------------
    # 4. 실물 보드 J2 기준 1:1 하드웨어 배선 및 전원 분배표
    # ---------------------------------------------------------
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(6)
    r_h3 = h3.add_run("3. 실물 디스플레이 보드 (J2) 1:1 하드웨어 배선 및 전원 분배표")
    r_h3.font.size = Pt(13)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A)

    table3_data = [
        ("ESP32 J2 실크 라벨", "J2 물리 위치", "연결 대상 부품", "부품 측 핀", "배선 색상", "브레드보드 경유 여부 및 상세 가이드"),
        (
            "VBUS",
            "상단 J2 맨 좌측\n(1번 구멍)",
            "전원 공통 레일",
            "VCC",
            "빨강 (Red)",
            "브레드보드 [+] 전원 레일에 직결 ➔ 세그먼트 VCC 및 HX711 VCC로 분배 공급"
        ),
        (
            "GND",
            "상단 J2 좌측 2번째\n(2번 구멍)",
            "접지 공통 레일",
            "GND",
            "검정 (Black)",
            "브레드보드 [-] 접지 레일에 직결 ➔ 세그먼트 GND 및 HX711 GND로 공통 접지"
        ),
        (
            "17",
            "상단 J2 좌측 13번째\n(GPIO 17)",
            "MAX7219 세그먼트",
            "DIN",
            "노랑 / 파랑",
            "시리얼 데이터 라인. 점퍼선으로 세그먼트 DIN 단자에 1:1 직결"
        ),
        (
            "18",
            "상단 J2 좌측 8번째\n(GPIO 18)",
            "MAX7219 세그먼트",
            "CLK",
            "초록 / 흰색",
            "시리얼 클록 라인. 점퍼선으로 세그먼트 CLK 단자에 1:1 직결"
        ),
        (
            "19",
            "상단 J2 좌측 5번째\n(GPIO 19)",
            "MAX7219 세그먼트",
            "CS (LOAD)",
            "주황 / 보라",
            "칩 셀렉트 라인. 점퍼선으로 세그먼트 CS 단자에 1:1 직결"
        ),
        (
            "TX\n(GPIO 43)",
            "상단 J2 좌측 4번째\n(자유 GPIO)",
            "HX711 로드셀 앰프",
            "DT (DOUT)",
            "파랑",
            "⭐️ [권장] 24비트 하중 데이터 수신 라인. LCD 백라이트 충돌 없는 완전 자유 핀"
        ),
        (
            "RX\n(GPIO 44)",
            "상단 J2 좌측 3번째\n(자유 GPIO)",
            "HX711 로드셀 앰프",
            "SCK",
            "초록",
            "⭐️ [권장] 로드셀 클록 펄스 라인. LCD 컬러 버스와 충돌 없는 완전 자유 핀"
        ),
        (
            "IO10\n(선택)",
            "상단 J2 좌측 12번째\n(GPIO 10)",
            "패시브 부저",
            "I/O (+)",
            "갈색 / 노랑",
            "피에조 부저 신호 핀. (-)는 브레드보드 GND 레일에 연결"
        )
    ]

    t3 = doc.add_table(rows=len(table3_data), cols=6)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths3 = [Inches(1.1), Inches(1.2), Inches(1.3), Inches(0.8), Inches(0.8), Inches(1.7)]

    for row_idx, row in enumerate(t3.rows):
        data_row = table3_data[row_idx]
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths3[col_idx]
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            if row_idx == 0:
                set_cell_background(cell, "0F294A")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(data_row[col_idx])
                run.font.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if col_idx in [0, 1, 3, 4]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(data_row[col_idx])
                run.font.size = Pt(8.5)
                if col_idx == 0:
                    run.font.bold = True

    # ---------------------------------------------------------
    # 5. 핵심 제작 노하우 & 주의사항
    # ---------------------------------------------------------
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(8)
    h4.paragraph_format.space_after = Pt(6)
    r_h4 = h4.add_run("4. 하드웨어 제작 & 납땜 필수 주의사항 (체크리스트)")
    r_h4.font.size = Pt(13)
    r_h4.font.bold = True
    r_h4.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A)

    tips = [
        ("⚠️ 로드셀 HX711 핀 충돌 방지: ", "기존 소스코드의 GPIO 2는 LCD 백라이트 핀(LCD-BL-EN)이고, GPIO 3은 LCD 파란색 신호선(LCD-B1)입니다. 로드셀을 이 핀에 연결하면 화면 백라이트가 꺼지거나 그래픽이 깨지므로, 반드시 위 배선표대로 J2에 있는 완전 자유 핀인 'TX(GPIO 43)'와 'RX(GPIO 44)'를 사용해야 합니다."),
        ("⚠️ MAX7219 세그먼트 입력단(IN) 확인: ", "세그먼트 기판 양쪽 중 'DIN' 단자가 있는 입력 헤더에 연결해야 합니다. 반대쪽 'DOUT'은 다중 모듈 연장용이므로 연결 시 동작하지 않습니다."),
        ("💡 핀헤더 수(Male) vs 암(Female) 선택: ", "빠른 작업과 결합력을 원하시면 '수(Male) 핀헤더 + 암-암 점퍼선'을 추천하며, 합선 방지와 슬림한 액자 프레임 매립을 원하시면 '암(Female) 핀소켓 + 암-수 점퍼선'을 추천합니다."),
        ("💡 브레드보드 전원 레일 공유: ", "디스플레이 보드의 VBUS(5V)와 GND를 먼저 브레드보드의 긴 전원 버스 라인(+, -)에 연결하고, 세그먼트와 HX711의 전원을 여기서 나눠서 가져가면 배선이 꼬이지 않고 매우 깔끔해집니다.")
    ]

    for title, body in tips:
        tp = doc.add_paragraph()
        tp.paragraph_format.space_before = Pt(3)
        tp.paragraph_format.space_after = Pt(3)
        rt = tp.add_run("• " + title)
        rt.font.bold = True
        rt.font.size = Pt(9.5)
        if "⚠️" in title:
            rt.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
        else:
            rt.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)
        rb = tp.add_run(body)
        rb.font.size = Pt(9.5)

    # 저장 경로
    save_path = r"c:\Users\user\Desktop\기존 바탕화면 자료\바탕화면 폴더 정리\jeayonglee\machineartlab\박건작가님\부품리스트.docx"
    doc.save(save_path)
    print(f"Document saved successfully: {save_path}")

    # 백업 겸 추가 파일로도 저장
    backup_path = r"c:\Users\user\Desktop\기존 바탕화면 자료\바탕화면 폴더 정리\jeayonglee\machineartlab\박건작가님\머신아트랩_박건작가님_부품리스트_및_배선가이드.docx"
    doc.save(backup_path)
    print(f"Backup saved successfully: {backup_path}")

if __name__ == "__main__":
    create_document()
